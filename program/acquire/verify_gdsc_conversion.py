from openpyxl import load_workbook
from pathlib import Path
import csv,json,math,itertools
p=Path('data/raw/GDSC/release8.5/GDSC2_fitted_dose_response_27Oct23.xlsx');w=load_workbook(p,read_only=True,data_only=True)
print('sheets',w.sheetnames,flush=True);ws=w[w.sheetnames[0]];it=ws.iter_rows(values_only=True);header=list(next(it));mismatch=[];total=0;max_abs=0;numeric=0; mismatch_count=0; whitespace_only=0
with open('data/canonical/GDSC/release8.5/GDSC2_fitted_dose_response_27Oct23.csv',encoding='utf-8-sig',newline='') as f:
 reader=csv.reader(f);ch=next(reader);same=ch==header
 for line,(xr,cr) in enumerate(itertools.zip_longest(it,reader),start=2):
  total+=1
  if xr is None or cr is None:
   mismatch.append({'line':line,'reason':'row_count'});continue
  if len(xr)!=len(cr):mismatch.append({'line':line,'reason':'width'});continue
  for j,(a,b) in enumerate(zip(xr,cr)):
   if isinstance(a,(float,int)):
    try:
     v=float(b);diff=abs(a-v);max_abs=max(max_abs,diff);numeric+=1
     ok=math.isclose(a,v,rel_tol=1e-12,abs_tol=1e-12)
    except ValueError:ok=False
   else:ok=('' if a is None else str(a))==b
   if not ok:
    mismatch_count+=1
    if ('' if a is None else str(a)).strip()==b.strip():whitespace_only+=1
    if len(mismatch)<5:mismatch.append({'line':line,'field':header[j],'xlsx':a,'csv':b})
r={'xlsx':p.as_posix(),'csv':'data/canonical/GDSC/release8.5/GDSC2_fitted_dose_response_27Oct23.csv','rows_compared':total,'header_equal':same,'numeric_cells_compared':numeric,'max_numeric_absolute_difference':max_abs,'mismatch_count':mismatch_count,'whitespace_only_mismatches':whitespace_only,'mismatch_examples':mismatch,'comparison':'numeric tolerance rel=1e-12 abs=1e-12; string exact; no rounding changes to source'}
Path('data/manifests/gdsc_xlsx_csv_comparison.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(r,ensure_ascii=False))
